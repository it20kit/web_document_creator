from docx.shared import *
from document_creator.document_block_handlers.blocks_handlers.abstract_block_creator import AbstractBlockCreator
from docx import Document


class TableCreator(AbstractBlockCreator):
    """Умеет создавать различные таблицы"""

    def create_block_based_on_data(self, document: Document, data: dict) -> None:
        """Создает таблицу с любым количеством столбцов и строк по переданным данным"""
        content = self.document_structure_manager.get_content_from_block(data)
        styles = self.document_structure_manager.get_style_from_block(data)
        size = self.document_structure_manager.get_table_size_from_block(data)
        rows = self.__count_number_of_rows(content)
        cols = self.__count_number_of_cols(content)
        table = self.__create_empty_table_by_parameters(document, rows, cols)
        self.__set_size_of_table(table, size)
        for number_row in range(len(content)):
            text_by_cell = self.get_value_from_content(content[number_row])
            self.__add_data_to_the_table(table, text_by_cell, styles, number_row)

    def __add_data_to_the_table(self, table, content: list, style: list, number_row: int) -> None:
        """Заполнить таблицу данными"""
        for number_cols in range(len(content)):
            text = content[number_cols]
            if self.is_value_empty(text) is False:
                cell = table.cell(number_row, number_cols)
                if number_cols < len(style):
                    style_by_cell = style[number_cols]
                else:
                    style_by_cell = {}
                paragraph_in_cells = cell.paragraphs[0]
                self.add_content_in_paragraph(paragraph_in_cells, text, style_by_cell)

    @staticmethod
    def __create_empty_table_by_parameters(document: Document, rows: int, cols: int):
        """Создает пустую таблицу по числу строк и столбцов"""
        table = document.add_table(rows=rows, cols=cols)
        table.style = 'Table Grid'
        table.width = 1
        return table

    @staticmethod
    def __count_number_of_rows(content: list) -> int:
        """Считает сколько строк в таблице по полученным данным"""
        return len(content)

    def __count_number_of_cols(self, content: list) -> int:
        """Считает сколько столбцов в таблице по полученным данным"""
        return len(self.get_value_from_content(content[0]))

    @staticmethod
    def __set_size_of_table(table, sizes: list) -> None:
        """Установить размеры ячеек для всей таблицы по переданным данным"""
        for i in range(len(sizes)):
            table.columns[i].width = Cm(sizes[i])
